"""analyze_bottlenecks and save_excel_report, checked against hand-computed numbers."""
import os
import sys
from datetime import date, timedelta

import openpyxl
import pandas as pd
import pytest
from fastapi.testclient import TestClient

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))
import app as app_module  # noqa: E402
import tools  # noqa: E402

T = date.today()


def d(days_ago):
    return T - timedelta(days=days_ago)


ROWS = [
    # id, task, dept, owner, deadline, status, priority
    ("T1", "Payroll", "HR", "Priya Sharma", d(10), "Pending", "High"),
    ("T2", "Handbook", "HR", "Priya Sharma", d(2), "Blocked", "Low"),
    ("T3", "Proposal", "Sales", "Ankit Desai", d(5), "Blocked", "High"),
    ("T4", "Forecast", "Sales", "Ankit Desai", T + timedelta(days=4), "Pending", "Medium"),
    ("T5", "Patching", "IT", "Vikram Rao", d(20), "WIP", "Medium"),
    ("T6", "Old but done", "IT", "Vikram Rao", d(30), "done", "High"),  # completed: must be ignored
    ("T7", "No date", "IT", "Vikram Rao", None, "Pending", "Low"),
    ("T8", "Odd priority", "Ops", "Meera Joshi", d(1), "Pending", "URGENT"),
]


@pytest.fixture(autouse=True)
def tracker(tmp_path, monkeypatch):
    path = tmp_path / "tasks.xlsx"
    pd.DataFrame(ROWS, columns=tools.COLUMNS).to_excel(path, index=False)
    monkeypatch.setenv("TASKS_FILE", str(path))
    monkeypatch.chdir(tmp_path)
    return tmp_path


def test_workload_per_owner_ignores_completed_tasks():
    w = {x["owner"]: x for x in tools.analyze_bottlenecks()["workload_by_owner"]}
    assert w["Priya Sharma"] == {"owner": "Priya Sharma", "open_tasks": 2, "overdue": 2, "high_priority_open": 1, "blocked": 1}
    assert w["Vikram Rao"]["open_tasks"] == 2 and w["Vikram Rao"]["overdue"] == 1  # T6 completed, T7 no deadline
    assert w["Ankit Desai"]["open_tasks"] == 2 and w["Ankit Desai"]["overdue"] == 1


def test_highlights_name_the_right_owner_and_department():
    h = tools.analyze_bottlenecks()["highlights"]
    assert h["most_overloaded_owner"] == "Priya Sharma"  # 2 overdue vs 1 for everyone else
    assert h["slowest_department"] == "IT"  # only T5, 20 days late


def test_delay_per_department():
    delay = {x["department"]: x for x in tools.analyze_bottlenecks()["delay_by_department"]}
    assert delay["HR"]["avg_days_overdue"] == 6.0 and delay["HR"]["max_days_overdue"] == 10  # (10 + 2) / 2
    assert delay["IT"]["avg_days_overdue"] == 20.0 and delay["IT"]["overdue"] == 1
    assert delay["Sales"]["avg_days_overdue"] == 5.0
    assert delay["Ops"]["avg_days_overdue"] == 1.0


def test_overdue_by_priority_counts_unknown_priority():
    assert tools.analyze_bottlenecks()["overdue_by_priority"] == {"High": 2, "Medium": 1, "Low": 1, "Unknown": 1}


def test_blocked_tasks_sorted_high_priority_first():
    assert [t["task_id"] for t in tools.analyze_bottlenecks()["blocked_tasks"]] == ["T3", "T2"]


def test_needs_attention_first_ranks_priority_then_lateness():
    r = tools.analyze_bottlenecks()
    assert [t["task_id"] for t in r["needs_attention_first"]] == ["T1", "T3", "T5", "T2", "T8"]
    assert r["needs_attention_first"][0]["days_overdue"] == 10
    assert [t["task_id"] for t in tools.analyze_bottlenecks(top_n=2)["needs_attention_first"]] == ["T1", "T3"]


def test_missing_deadline_reported_and_bad_top_n_does_not_crash():
    r = tools.analyze_bottlenecks(top_n="abc")
    assert r["open_tasks_without_deadline"] == 1
    assert len(r["needs_attention_first"]) == 5


def test_error_when_file_missing(monkeypatch):
    monkeypatch.setenv("TASKS_FILE", "nope.xlsx")
    assert "error" in tools.analyze_bottlenecks()
    assert "error" in tools.save_excel_report()


def test_excel_report_contents(tracker):
    result = tools.save_excel_report("Weekly Ops", "HR is slipping. Priya needs help.")
    assert result["saved_to"].endswith(".xlsx") and os.path.exists(result["saved_to"])
    wb = openpyxl.load_workbook(result["saved_to"])
    assert wb.sheetnames == ["Summary", "Departments", "Overdue", "Owner workload", "Blocked", "Delay by department"]

    summary = {r[0].value: r[1].value for r in wb["Summary"].iter_rows(min_row=5, max_row=11)}
    assert summary["Total tasks"] == 8 and summary["Completed"] == 1 and summary["Open"] == 7
    assert summary["Overdue"] == 5 and summary["Most overloaded owner"] == "Priya Sharma"
    assert any("Priya needs help" in str(c.value) for row in wb["Summary"].iter_rows() for c in row if c.value)

    overdue = list(wb["Overdue"].iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in overdue] == ["T1", "T3", "T5", "T2", "T8"]  # priority, then days late
    assert isinstance(overdue[0][4], (date, pd.Timestamp)) or hasattr(overdue[0][4], "year")  # real Excel date
    assert overdue[0][7] == 10
    assert wb["Overdue"].freeze_panes == "A2"
    assert wb["Departments"].max_row == 5  # header + HR, IT, Ops, Sales


def test_excel_report_without_narrative_has_no_commentary_block():
    result = tools.save_excel_report("Plain")
    wb = openpyxl.load_workbook(result["saved_to"])
    assert not any(c.value == "Commentary" for row in wb["Summary"].iter_rows() for c in row)


def test_excel_report_locked_file_gives_friendly_error(monkeypatch):
    def locked(self, *a, **k):
        raise PermissionError("in use")

    monkeypatch.setattr(openpyxl.workbook.workbook.Workbook, "save", locked)
    assert "open in Excel" in tools.save_excel_report()["error"]


def test_api_and_dashboard_page_expose_bottlenecks():
    c = TestClient(app_module.app)
    r = c.get("/api/bottlenecks")
    assert r.status_code == 200 and r.json()["highlights"]["slowest_department"] == "IT"
    assert "loadBottlenecks" in open(os.path.join(os.path.dirname(app_module.__file__), "static", "index.html"), encoding="utf-8").read()


def test_overloaded_owner_ties_break_on_high_priority_work(tmp_path, monkeypatch):
    rows = [
        ("A1", "x", "HR", "Aaron Low", d(3), "Pending", "Low"),  # 1 overdue, 0 high
        ("B1", "y", "HR", "Zoe High", d(3), "Pending", "High"),  # 1 overdue, 1 high -> should win the tie
    ]
    path = tmp_path / "tie.xlsx"
    pd.DataFrame(rows, columns=tools.COLUMNS).to_excel(path, index=False)
    monkeypatch.setenv("TASKS_FILE", str(path))
    assert tools.analyze_bottlenecks()["highlights"]["most_overloaded_owner"] == "Zoe High"
