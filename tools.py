"""Tools the agent can call.

Every tool returns plain JSON-serialisable dicts so results can be fed back
to the LLM. Tools never raise to the agent for expected problems; they return
{"error": "..."} so the model can read the message and choose another step.
"""
import json
import os
import shutil
import zipfile
from datetime import date, datetime

import pandas as pd

COLUMNS = ["task_id", "task", "department", "owner", "deadline", "status", "priority"]
VALID_STATUSES = ["Pending", "In Progress", "Blocked", "Completed"]
STATUS_ALIASES = {
    "pending": "Pending",
    "todo": "Pending",
    "to do": "Pending",
    "not started": "Pending",
    "in progress": "In Progress",
    "inprogress": "In Progress",
    "wip": "In Progress",
    "ongoing": "In Progress",
    "blocked": "Blocked",
    "on hold": "Blocked",
    "completed": "Completed",
    "complete": "Completed",
    "done": "Completed",
    "finished": "Completed",
}
VALID_PRIORITIES = ["High", "Medium", "Low"]


SOURCE_FILE = "data/source.json"


def tasks_path() -> str:
    """Active tracker: the file chosen in the app, else TASKS_FILE, else data/tasks.xlsx."""
    try:
        with open(SOURCE_FILE, encoding="utf-8") as f:
            chosen = json.load(f).get("path")
        if chosen:
            return chosen
    except (OSError, ValueError):
        pass
    return os.getenv("TASKS_FILE", "data/tasks.xlsx")


def set_tasks_path(path: str) -> dict:
    """Point the app at any .xlsx/.csv tracker on this computer (validated first)."""
    path = (path or "").strip().strip('"').strip("'")
    if not path.lower().endswith((".xlsx", ".csv")):
        return {"error": "Choose an .xlsx or .csv file"}
    path = os.path.abspath(os.path.expanduser(path))
    try:
        df = load_tasks(path)
    except ValueError as e:
        return {"error": str(e)}
    os.makedirs(os.path.dirname(SOURCE_FILE), exist_ok=True)
    with open(SOURCE_FILE, "w", encoding="utf-8") as f:
        json.dump({"path": path}, f)
    return {"path": path, "rows": int(len(df))}


def reset_tasks_path() -> None:
    try:
        os.remove(SOURCE_FILE)
    except OSError:
        pass


def file_status() -> dict:
    """Cheap check the dashboard polls to notice when the Excel file is saved."""
    path = tasks_path()
    try:
        st = os.stat(path)
    except OSError:
        return {"path": path, "exists": False}
    return {"path": path, "exists": True, "modified": st.st_mtime, "size": st.st_size}


def normalize_status(value) -> str:
    key = str(value).strip().lower()
    return STATUS_ALIASES.get(key, "Unknown")


def load_tasks(path: str | None = None) -> pd.DataFrame:
    """Load and clean the tracker. Raises ValueError for unusable files."""
    path = path or tasks_path()
    if not os.path.exists(path):
        raise ValueError(f"Task file not found: {path}")
    try:
        if path.lower().endswith(".csv"):
            df = pd.read_csv(path)
        else:
            df = pd.read_excel(path)
    except PermissionError:
        raise ValueError(f"Cannot read {path}: the file is locked. Save and close it, then try again.")
    except (zipfile.BadZipFile, pd.errors.EmptyDataError, KeyError):
        raise ValueError(f"Could not read {path}: it may be mid-save or not a valid file. Try again in a moment.")

    df.columns = [str(c).strip().lower().replace(" ", "_") for c in df.columns]
    missing = [c for c in COLUMNS if c not in df.columns]
    if missing:
        raise ValueError(f"Task file is missing required columns: {missing}")

    df = df[COLUMNS].copy()
    df["status"] = df["status"].apply(normalize_status)
    df["deadline"] = pd.to_datetime(df["deadline"], errors="coerce")
    df["priority"] = df["priority"].astype(str).str.strip().str.title()
    df["priority"] = df["priority"].where(df["priority"].isin(VALID_PRIORITIES), "Unknown")
    for col in ["task_id", "task", "department", "owner"]:
        df[col] = df[col].astype(str).str.strip()
    return df


def _today() -> date:
    return date.today()


def _records(df: pd.DataFrame) -> list[dict]:
    out = []
    for r in df.to_dict("records"):
        r = dict(r)
        d = r.get("deadline")
        r["deadline"] = None if pd.isna(d) else pd.Timestamp(d).date().isoformat()
        if "days_overdue" in r:
            r["days_overdue"] = int(r["days_overdue"])
        out.append(r)
    return out


# ---------------------------------------------------------------- read tools
def read_tasks(department: str | None = None, status: str | None = None, owner: str | None = None) -> dict:
    """Return tasks, optionally filtered."""
    try:
        df = load_tasks()
    except ValueError as e:
        return {"error": str(e)}
    if department:
        df = df[df["department"].str.lower() == department.strip().lower()]
    if status:
        wanted = normalize_status(status)
        if wanted == "Unknown":
            return {"error": f"Unknown status '{status}'. Valid: {VALID_STATUSES}"}
        df = df[df["status"] == wanted]
    if owner:
        df = df[df["owner"].str.lower() == owner.strip().lower()]
    return {"count": int(len(df)), "tasks": _records(df)}


def get_overdue_tasks(as_of: str | None = None) -> dict:
    """Open tasks whose deadline has passed. Also reports rows with unusable data."""
    try:
        df = load_tasks()
    except ValueError as e:
        return {"error": str(e)}
    if as_of:
        try:
            ref = datetime.strptime(as_of, "%Y-%m-%d").date()
        except ValueError:
            return {"error": "as_of must be formatted YYYY-MM-DD"}
    else:
        ref = _today()

    open_df = df[df["status"] != "Completed"]
    no_deadline = open_df[open_df["deadline"].isna()]
    dated = open_df.dropna(subset=["deadline"]).copy()
    dated["days_overdue"] = (pd.Timestamp(ref) - dated["deadline"]).dt.days
    overdue = dated[dated["days_overdue"] > 0].sort_values("days_overdue", ascending=False)
    return {
        "as_of": ref.isoformat(),
        "overdue_count": int(len(overdue)),
        "overdue_tasks": _records(overdue),
        "tasks_with_missing_deadline": _records(no_deadline),
        "tasks_with_unrecognised_status": int((df["status"] == "Unknown").sum()),
    }


def summarize_by_department() -> dict:
    """Status counts and overdue counts per department."""
    try:
        df = load_tasks()
    except ValueError as e:
        return {"error": str(e)}
    ref = pd.Timestamp(_today())
    summary = {}
    for dept, g in df.groupby("department"):
        open_g = g[g["status"] != "Completed"]
        overdue = open_g[open_g["deadline"].notna() & (open_g["deadline"] < ref)]
        summary[dept] = {
            "total": int(len(g)),
            "completed": int((g["status"] == "Completed").sum()),
            "in_progress": int((g["status"] == "In Progress").sum()),
            "pending": int((g["status"] == "Pending").sum()),
            "blocked": int((g["status"] == "Blocked").sum()),
            "overdue": int(len(overdue)),
        }
    return {"as_of": _today().isoformat(), "departments": summary}


# ------------------------------------------------------------ write / draft
def update_task_status(task_id: str, new_status: str) -> dict:
    """CONSEQUENTIAL: change a task's status in the tracker (needs approval)."""
    status = normalize_status(new_status)
    if status == "Unknown":
        return {"error": f"Invalid status '{new_status}'. Valid: {VALID_STATUSES}"}
    try:
        df = load_tasks()
    except ValueError as e:
        return {"error": str(e)}
    mask = df["task_id"] == str(task_id).strip()
    if not mask.any():
        return {"error": f"No task with id '{task_id}'"}

    path = tasks_path()
    root, ext = os.path.splitext(path)
    ext = ext or ".xlsx"
    os.makedirs("data/backups", exist_ok=True)
    backup = f"data/backups/tasks_{datetime.now():%Y%m%d_%H%M%S}{ext}"
    try:
        shutil.copy(path, backup)
    except OSError as e:
        return {"error": f"Could not back up {path}: {e}"}

    old = df.loc[mask, "status"].iloc[0]
    df.loc[mask, "status"] = status
    out = df.copy()
    out["deadline"] = out["deadline"].dt.date
    tmp = f"{root}.tmp{ext}"  # write beside the real file, then swap in
    try:
        if ext.lower() == ".csv":
            out.to_csv(tmp, index=False)
        else:
            out.to_excel(tmp, index=False)
        os.replace(tmp, path)
    except OSError as e:
        if os.path.exists(tmp):
            os.remove(tmp)
        if isinstance(e, PermissionError):
            return {
                "error": f"Could not write {path}: it is open in Excel. Save and close it, then press Approve again.",
                "retryable": True,
            }
        return {"error": f"Could not write {path}: {e}"}
    return {"task_id": task_id, "old_status": old, "new_status": status, "backup": backup}


def draft_notification(owner: str, task_ids: list[str], tone: str = "polite") -> dict:
    """Draft (never sends) a reminder message for an owner about specific tasks."""
    try:
        df = load_tasks()
    except ValueError as e:
        return {"error": str(e)}
    ids = [str(t).strip() for t in (task_ids or [])]
    rows = df[(df["owner"].str.lower() == owner.strip().lower()) & (df["task_id"].isin(ids))]
    if rows.empty:
        return {"error": f"No matching tasks for owner '{owner}' and ids {ids}"}
    lines = []
    for r in _records(rows):
        lines.append(f"- {r['task_id']}: {r['task']} (due {r['deadline'] or 'no date'}, {r['status']}, {r['priority']} priority)")
    opener = "Quick reminder" if tone == "polite" else "Action needed"
    body = (
        f"Hi {owner.split()[0]},\n\n{opener}: the following tasks need your attention:\n\n"
        + "\n".join(lines)
        + "\n\nPlease update the tracker or let me know if anything is blocking you.\n\nThanks,\nOperations"
    )
    return {"to": owner, "subject": f"Task follow-up: {len(rows)} item(s)", "body": body, "sent": False}


def save_report(title: str, markdown: str) -> dict:
    """Save a finished report as a Markdown file in reports/."""
    if not markdown or not markdown.strip():
        return {"error": "Report body is empty"}
    os.makedirs("reports", exist_ok=True)
    safe = "".join(c if c.isalnum() else "_" for c in title.lower())[:40].strip("_") or "report"
    path = f"reports/{safe}_{_today().isoformat()}.md"
    with open(path, "w", encoding="utf-8") as f:
        f.write(f"# {title}\n\n{markdown}\n")
    return {"saved_to": path}


def send_notification(owner: str, subject: str, body: str) -> dict:
    """CONSEQUENTIAL: deliver a message to an owner (needs approval).

    Delivery is simulated: the message is written to outbox/ and recorded in
    outbox/sent_log.jsonl. Swap the body of this function for SMTP / Slack /
    Teams to make it real; the approval gate stays the same.
    """
    if not owner or not subject or not body:
        return {"error": "owner, subject and body are all required"}
    try:
        df = load_tasks()
    except ValueError as e:
        return {"error": str(e)}
    if not (df["owner"].str.lower() == owner.strip().lower()).any():
        return {"error": f"'{owner}' is not an owner in the tracker"}
    os.makedirs("outbox", exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    safe = "".join(c if c.isalnum() else "_" for c in owner.lower())
    path = f"outbox/{stamp}_{safe}.txt"
    with open(path, "w", encoding="utf-8") as f:
        f.write(f"To: {owner}\nSubject: {subject}\n\n{body}\n")
    with open("outbox/sent_log.jsonl", "a", encoding="utf-8") as f:
        f.write(json.dumps({"ts": datetime.now().isoformat(timespec="seconds"), "to": owner, "subject": subject, "file": path}) + "\n")
    return {"delivered_to": owner, "subject": subject, "outbox_file": path}


PRIORITY_RANK = {"High": 0, "Medium": 1, "Low": 2, "Unknown": 3}


def _task_row(r: dict) -> dict:
    """Plain JSON row for one task (days_overdue may be None when there is no deadline)."""
    d = r.get("deadline")
    late = r.get("days_overdue")
    return {
        "task_id": r["task_id"],
        "task": r["task"],
        "department": r["department"],
        "owner": r["owner"],
        "deadline": None if pd.isna(d) else pd.Timestamp(d).date().isoformat(),
        "status": r["status"],
        "priority": r["priority"],
        "days_overdue": None if pd.isna(late) else int(late),
    }


def analyze_bottlenecks(top_n: int = 10) -> dict:
    """Where is work getting stuck? Workload per owner, blocked tasks, overdue by priority, delay per department."""
    try:
        df = load_tasks()
    except ValueError as e:
        return {"error": str(e)}
    try:
        top_n = max(1, min(int(top_n), 50))
    except (TypeError, ValueError):
        top_n = 10

    ref = pd.Timestamp(_today())
    open_df = df[df["status"] != "Completed"].copy()
    open_df["days_overdue"] = (ref - open_df["deadline"]).dt.days  # NaN when there is no deadline
    open_df["is_overdue"] = open_df["days_overdue"] > 0
    open_df["prio_rank"] = open_df["priority"].map(PRIORITY_RANK).fillna(3)
    overdue = open_df[open_df["is_overdue"]]

    workload = []
    for owner, g in open_df.groupby("owner"):
        workload.append(
            {
                "owner": owner,
                "open_tasks": int(len(g)),
                "overdue": int(g["is_overdue"].sum()),
                "high_priority_open": int((g["priority"] == "High").sum()),
                "blocked": int((g["status"] == "Blocked").sum()),
            }
        )
    workload.sort(key=lambda w: (-w["overdue"], -w["high_priority_open"], -w["open_tasks"], w["owner"]))

    blocked = open_df[open_df["status"] == "Blocked"].sort_values(["prio_rank", "deadline"], na_position="last")

    delay = []
    for dept, g in overdue.groupby("department"):
        delay.append(
            {
                "department": dept,
                "overdue": int(len(g)),
                "avg_days_overdue": round(float(g["days_overdue"].mean()), 1),
                "max_days_overdue": int(g["days_overdue"].max()),
            }
        )
    delay.sort(key=lambda d: (-d["avg_days_overdue"], d["department"]))

    attention = overdue.sort_values(["prio_rank", "days_overdue"], ascending=[True, False]).head(top_n)

    return {
        "as_of": _today().isoformat(),
        "highlights": {
            "most_overloaded_owner": workload[0]["owner"] if workload and workload[0]["overdue"] > 0 else None,
            "slowest_department": delay[0]["department"] if delay else None,
        },
        "workload_by_owner": workload,
        "blocked_tasks": [_task_row(r) for r in blocked.to_dict("records")],
        "overdue_by_priority": {p: int((overdue["priority"] == p).sum()) for p in ["High", "Medium", "Low", "Unknown"]},
        "delay_by_department": delay,
        "needs_attention_first": [_task_row(r) for r in attention.to_dict("records")],
        "open_tasks_without_deadline": int(open_df["deadline"].isna().sum()),
    }


def _write_sheet(wb, name: str, headers: list, rows: list, date_cols: tuple = ()):
    """One formatted worksheet: bold header, frozen top row, filter, sensible widths."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(name)
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F5BEA")
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append([date.fromisoformat(v) if (i in date_cols and isinstance(v, str)) else v for i, v in enumerate(r)])
    for i in date_cols:
        for (cell,) in ws.iter_rows(min_row=2, min_col=i + 1, max_col=i + 1):
            cell.number_format = "yyyy-mm-dd"
    for idx in range(1, len(headers) + 1):
        longest = max(len(str(c.value)) if c.value is not None else 0 for c in ws[get_column_letter(idx)])
        ws.column_dimensions[get_column_letter(idx)].width = min(50, max(12, longest + 2))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def save_excel_report(title: str = "Weekly Operations Report", narrative: str = "") -> dict:
    """Save an Excel workbook report. All numbers come from the tracker, not from the model;
    `narrative` is the model's written commentary and goes on the Summary sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    summary = summarize_by_department()
    overdue = get_overdue_tasks()
    bott = analyze_bottlenecks(top_n=50)
    for r in (summary, overdue, bott):
        if "error" in r:
            return {"error": r["error"]}

    depts = summary["departments"]
    total = sum(d["total"] for d in depts.values())
    done = sum(d["completed"] for d in depts.values())

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Generated {_today().isoformat()}"
    ws["A4"] = "Key figures"
    ws["A4"].font = Font(bold=True)
    figures = [
        ("Total tasks", total),
        ("Completed", done),
        ("Open", total - done),
        ("Overdue", overdue["overdue_count"]),
        ("Open tasks without a deadline", len(overdue["tasks_with_missing_deadline"])),
        ("Most overloaded owner", bott["highlights"]["most_overloaded_owner"] or "none"),
        ("Slowest department", bott["highlights"]["slowest_department"] or "none"),
    ]
    row = 5
    for label, value in figures:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=value)
        row += 1
    if narrative and narrative.strip():
        row += 1
        ws.cell(row=row, column=1, value="Commentary").font = Font(bold=True)
        row += 1
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        cell = ws.cell(row=row, column=1, value=narrative.strip())
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 16 * max(2, -(-len(narrative) // 100))
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 22
    for col in "CDEF":
        ws.column_dimensions[col].width = 16

    _write_sheet(
        wb,
        "Departments",
        ["Department", "Total", "Completed", "In progress", "Pending", "Blocked", "Overdue"],
        [[d, s["total"], s["completed"], s["in_progress"], s["pending"], s["blocked"], s["overdue"]] for d, s in depts.items()],
    )
    ordered = sorted(overdue["overdue_tasks"], key=lambda t: (PRIORITY_RANK.get(t["priority"], 3), -t["days_overdue"]))
    task_cols = ["Task ID", "Task", "Department", "Owner", "Deadline", "Status", "Priority", "Days overdue"]
    _write_sheet(
        wb,
        "Overdue",
        task_cols,
        [[t["task_id"], t["task"], t["department"], t["owner"], t["deadline"], t["status"], t["priority"], t["days_overdue"]] for t in ordered],
        date_cols=(4,),
    )
    _write_sheet(
        wb,
        "Owner workload",
        ["Owner", "Open tasks", "Overdue", "High priority open", "Blocked"],
        [[w["owner"], w["open_tasks"], w["overdue"], w["high_priority_open"], w["blocked"]] for w in bott["workload_by_owner"]],
    )
    _write_sheet(
        wb,
        "Blocked",
        task_cols,
        [[t["task_id"], t["task"], t["department"], t["owner"], t["deadline"], t["status"], t["priority"], t["days_overdue"]] for t in bott["blocked_tasks"]],
        date_cols=(4,),
    )
    _write_sheet(
        wb,
        "Delay by department",
        ["Department", "Overdue tasks", "Average days overdue", "Max days overdue"],
        [[d["department"], d["overdue"], d["avg_days_overdue"], d["max_days_overdue"]] for d in bott["delay_by_department"]],
    )

    os.makedirs("reports", exist_ok=True)
    safe = "".join(c if c.isalnum() else "_" for c in title.lower())[:40].strip("_") or "report"
    path = f"reports/{safe}_{_today().isoformat()}.xlsx"
    try:
        wb.save(path)
    except PermissionError:
        return {"error": f"Could not save {path}: it is open in Excel. Close it and try again."}
    except OSError as e:
        return {"error": f"Could not save {path}: {e}"}
    return {"saved_to": path, "sheets": wb.sheetnames, "overdue_rows": len(ordered)}


# ------------------------------------------------------------------ registry
TOOLS = {
    "analyze_bottlenecks": {
        "fn": analyze_bottlenecks,
        "requires_approval": False,
        "description": "Find where work is getting stuck: open workload per owner, blocked tasks, overdue counts by priority, average/max delay per department, and a ranked list of overdue tasks to handle first.",
        "parameters": {
            "type": "OBJECT",
            "properties": {"top_n": {"type": "INTEGER", "description": "How many tasks to put in the needs_attention_first list (default 10)"}},
        },
    },
    "save_excel_report": {
        "fn": save_excel_report,
        "requires_approval": False,
        "description": "Save the weekly report as a formatted Excel workbook (Summary, Departments, Overdue, Owner workload, Blocked, Delay by department). All figures are computed from the tracker; pass only a short written commentary as `narrative`.",
        "parameters": {
            "type": "OBJECT",
            "properties": {"title": {"type": "STRING"}, "narrative": {"type": "STRING", "description": "Short commentary for the Summary sheet"}},
        },
    },
    "send_notification": {
        "fn": send_notification,
        "requires_approval": True,
        "description": "Send a message to a task owner (usually the text produced by draft_notification). Queued for human approval; NOT sent immediately.",
        "parameters": {
            "type": "OBJECT",
            "properties": {
                "owner": {"type": "STRING"},
                "subject": {"type": "STRING"},
                "body": {"type": "STRING"},
            },
            "required": ["owner", "subject", "body"],
        },
    },
    "read_tasks": {
        "fn": read_tasks,
        "requires_approval": False,
        "description": "Read tasks from the tracker. Optional filters: department, status, owner.",
        "parameters": {
            "type": "OBJECT",
            "properties": {
                "department": {"type": "STRING"},
                "status": {"type": "STRING", "description": "Pending, In Progress, Blocked or Completed"},
                "owner": {"type": "STRING"},
            },
        },
    },
    "get_overdue_tasks": {
        "fn": get_overdue_tasks,
        "requires_approval": False,
        "description": "List open tasks past their deadline with days overdue, plus tasks with missing deadlines.",
        "parameters": {
            "type": "OBJECT",
            "properties": {"as_of": {"type": "STRING", "description": "Optional date YYYY-MM-DD, default today"}},
        },
    },
    "summarize_by_department": {
        "fn": summarize_by_department,
        "requires_approval": False,
        "description": "Per-department counts of total, completed, in-progress, pending, blocked and overdue tasks.",
        "parameters": {"type": "OBJECT", "properties": {}},
    },
    "draft_notification": {
        "fn": draft_notification,
        "requires_approval": False,
        "description": "Draft (not send) a reminder message to a task owner about specific task ids.",
        "parameters": {
            "type": "OBJECT",
            "properties": {
                "owner": {"type": "STRING"},
                "task_ids": {"type": "ARRAY", "items": {"type": "STRING"}},
                "tone": {"type": "STRING", "description": "polite or urgent"},
            },
            "required": ["owner", "task_ids"],
        },
    },
    "save_report": {
        "fn": save_report,
        "requires_approval": False,
        "description": "Save the final weekly report as a Markdown file.",
        "parameters": {
            "type": "OBJECT",
            "properties": {"title": {"type": "STRING"}, "markdown": {"type": "STRING"}},
            "required": ["title", "markdown"],
        },
    },
    "update_task_status": {
        "fn": update_task_status,
        "requires_approval": True,
        "description": "Change a task's status. This modifies the tracker, so it is queued for human approval and NOT executed immediately.",
        "parameters": {
            "type": "OBJECT",
            "properties": {
                "task_id": {"type": "STRING"},
                "new_status": {"type": "STRING", "description": "Pending, In Progress, Blocked or Completed"},
            },
            "required": ["task_id", "new_status"],
        },
    },
}
